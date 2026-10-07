from fastapi import APIRouter, Depends, HTTPException, status, Request
from fastapi.security import OAuth2PasswordRequestForm
from sqlalchemy.orm import Session
from datetime import timedelta, datetime
import os

from database import get_db
from models import User
from schemas import UserLogin, Token, UserCreate, UserResponse, ForgotPasswordRequest, ResetPasswordRequest
from auth import (
    authenticate_user,
    create_access_token,
    get_password_hash,
    ACCESS_TOKEN_EXPIRE_MINUTES,
    get_current_active_user
)
from audit_utils import log_login, log_logout, log_password_change, log_password_reset, get_client_info
from services.notification_service import NotificationService

router = APIRouter(prefix="/auth", tags=["Authentication"])

@router.post("/login", response_model=Token)
async def login(
    request: Request,
    form_data: OAuth2PasswordRequestForm = Depends(),
    db: Session = Depends(get_db)
):
    """
    Login endpoint to authenticate users and return JWT token
    """
    from models import LoginAttempt
    
    # Get client IP
    client_ip = request.client.host if request.client else "unknown"
    
    # Check if account is locked
    lockout_minutes = int(os.getenv("ACCOUNT_LOCKOUT_MINUTES", "15"))
    max_attempts = int(os.getenv("MAX_LOGIN_ATTEMPTS", "5"))
    lockout_time = datetime.now() - timedelta(minutes=lockout_minutes)
    
    recent_failed = db.query(LoginAttempt).filter(
        LoginAttempt.username == form_data.username,
        LoginAttempt.success == False,
        LoginAttempt.created_at > lockout_time
    ).count()
    
    if recent_failed >= max_attempts:
        # Get time of first failed attempt in lockout window
        first_attempt = db.query(LoginAttempt).filter(
            LoginAttempt.username == form_data.username,
            LoginAttempt.success == False,
            LoginAttempt.created_at > lockout_time
        ).order_by(LoginAttempt.created_at).first()
        
        if first_attempt:
            unlock_time = first_attempt.created_at + timedelta(minutes=lockout_minutes)
            minutes_remaining = int((unlock_time - datetime.now()).total_seconds() / 60) + 1
            raise HTTPException(
                status_code=status.HTTP_429_TOO_MANY_REQUESTS,
                detail=f"Account locked due to too many failed login attempts. Try again in {minutes_remaining} minutes."
            )
    
    user = authenticate_user(db, form_data.username, form_data.password)
    if not user:
        # Log failed attempt
        failed_attempt = LoginAttempt(
            username=form_data.username,
            ip_address=client_ip,
            success=False
        )
        db.add(failed_attempt)
        db.commit()
        
        # Calculate remaining attempts
        remaining = max_attempts - (recent_failed + 1)
        if remaining > 0:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=f"Incorrect username or password. {remaining} attempts remaining before account lockout.",
                headers={"WWW-Authenticate": "Bearer"},
            )
        else:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Incorrect username or password. Account will be locked for 15 minutes.",
                headers={"WWW-Authenticate": "Bearer"},
            )
    
    if not user.is_active:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="User account is inactive"
        )
    
    # Log successful attempt and clear failed attempts
    success_attempt = LoginAttempt(
        username=form_data.username,
        ip_address=client_ip,
        success=True
    )
    db.add(success_attempt)
    
    # Delete old failed attempts for this user
    db.query(LoginAttempt).filter(
        LoginAttempt.username == form_data.username,
        LoginAttempt.success == False
    ).delete()
    db.commit()
    
    # Audit log: successful login
    user_agent = request.headers.get("user-agent", "unknown")
    log_login(db, user.id, True, client_ip, user_agent)
    
    access_token_expires = timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    access_token = create_access_token(
        data={"sub": user.username, "role": user.role},
        expires_delta=access_token_expires
    )
    
    return {
        "access_token": access_token,
        "token_type": "bearer"
    }

@router.get("/me")
async def get_current_user_info(
    current_user: User = Depends(get_current_active_user),
    db: Session = Depends(get_db)
):
    """
    Get current logged-in user information with entity names
    """
    from models import District, Block, SubCentre, USGCentre, DeliveryPoint
    
    # Get entity names
    district_name = None
    block_name = None
    sub_centre_name = None
    usg_centre_name = None
    dp_name = None
    
    if current_user.district_id:
        district = db.query(District).filter(District.id == current_user.district_id).first()
        district_name = district.name if district else None
    
    if current_user.block_id:
        block = db.query(Block).filter(Block.id == current_user.block_id).first()
        block_name = block.name if block else None
    
    if current_user.sub_centre_id:
        sub_centre = db.query(SubCentre).filter(SubCentre.id == current_user.sub_centre_id).first()
        sub_centre_name = sub_centre.name if sub_centre else None
    
    if current_user.usg_centre_id:
        usg_centre = db.query(USGCentre).filter(USGCentre.id == current_user.usg_centre_id).first()
        usg_centre_name = usg_centre.name if usg_centre else None

    if current_user.dp_id:
        dp = db.query(DeliveryPoint).filter(DeliveryPoint.id == current_user.dp_id).first()
        dp_name = dp.name if dp else None
    
    return {
        "id": current_user.id,
        "username": current_user.username,
        "email": current_user.email,
        "role": current_user.role,
        "full_name": current_user.full_name,
        "mobile_number": current_user.mobile_number,
        "district_id": current_user.district_id,
        "district_name": district_name,
        "block_id": current_user.block_id,
        "block_name": block_name,
        "sub_centre_id": current_user.sub_centre_id,
        "sub_centre_name": sub_centre_name,
        "usg_centre_id": current_user.usg_centre_id,
        "usg_centre_name": usg_centre_name,
        "dp_id": current_user.dp_id,
        "dp_name": dp_name,
        "is_active": current_user.is_active,
        "created_at": current_user.created_at,
        "updated_at": current_user.updated_at
    }

@router.get("/roles")
async def get_available_roles():
    """
    Get list of available user roles for registration
    """
    return {
        "roles": [
            {"value": "district", "label": "District User"},
            {"value": "block", "label": "Block User"},
            {"value": "sub_centre", "label": "Sub-Centre User"},
            {"value": "usg_centre", "label": "USG Centre User"},
            {"value": "dp", "label": "Delivery Point User"},
            {"value": "pmsma", "label": "PMSMA User"}
        ]
    }

@router.get("/delivery-types")
async def get_delivery_types():
    """
    Get list of delivery type options for dropdown (no auth required)
    """
    from models import DeliveryType
    return [
        {"value": dt.value, "label": dt.value.replace("_", " ").title()}
        for dt in DeliveryType
    ]

@router.get("/password-requirements")
async def get_password_requirements():
    """
    Get password requirements for frontend validation
    """
    return {
        "requirements": [
            {"rule": "min_length", "value": 8, "description": "At least 8 characters"},
            {"rule": "uppercase", "value": 1, "description": "At least one uppercase letter (A-Z)"},
            {"rule": "lowercase", "value": 1, "description": "At least one lowercase letter (a-z)"},
            {"rule": "number", "value": 1, "description": "At least one number (0-9)"},
            {"rule": "special", "value": 1, "description": "At least one special character (!@#$%^&*)"},
        ],
        "example": "Example: MyP@ssw0rd"
    }

@router.get("/abnormal-findings-options")
async def get_abnormal_findings_options():
    """
    Get available abnormal findings options by trimester
    """
    import os
    from dotenv import load_dotenv
    load_dotenv()
    
    first_trimester = os.getenv("ABNORMAL_FINDINGS_FIRST_TRIMESTER", "").split(",")
    second_trimester = os.getenv("ABNORMAL_FINDINGS_SECOND_TRIMESTER", "").split(",")
    third_trimester = os.getenv("ABNORMAL_FINDINGS_THIRD_TRIMESTER", "").split(",")
    
    return {
        "first_trimester": [option.strip() for option in first_trimester if option.strip()],
        "second_trimester": [option.strip() for option in second_trimester if option.strip()],
        "third_trimester": [option.strip() for option in third_trimester if option.strip()]
    }

@router.get("/scan-types-options")
async def get_scan_types_options():
    """
    Get available scan types options for USG appointments
    """
    import os
    from dotenv import load_dotenv
    load_dotenv()
    
    scan_types = os.getenv("SCAN_TYPES", "").split(",")
    
    return {
        "scan_types": [option.strip() for option in scan_types if option.strip()]
    }

@router.get("/bulk-upload-template")
async def get_bulk_upload_template():
    """
    Get Excel template structure for pregnant women bulk upload
    """
    return {
        "template_columns": [
            "full_name",
            "mobile_number",
            # "aadhaar_number",
            "abha_id",
            "rch_id",
            "husband_name",
            "age",
            "ward_id",
            "block_id",
            "district_id",
            "sub_centre_id",
            "date_of_birth",
            "lmp_date",
            "edd_date",
            "gravida",
            "para",
            "blood_group",
            "address",
            "hpr_id"
        ],
        "required_columns": [
            "full_name",
            "mobile_number",
            "husband_name",
            "age",
            "ward_id"
        ],
        "sample_data": {
            "full_name": "Jane Doe",
            "mobile_number": "9876543210",
            # "aadhaar_number": "123456789012",
            "abha_id": "12-3456-7890-1234",
            "rch_id": "RCH123456",
            "husband_name": "John Doe",
            "age": 25,
            "ward_id": 1,
            "block_id": 1,
            "district_id": 1,
            "sub_centre_id": "(Optional - auto-assigned if ward mapping exists)",
            "date_of_birth": "1999-01-01",
            "lmp_date": "2024-01-01",
            "edd_date": "2024-10-08",
            "gravida": 2,
            "para": 1,
            "blood_group": "O+",
            "address": "123 Main Street",
            "hpr_id": "HPR123456"
        }
    }

@router.get("/download-bulk-upload-template")
async def download_bulk_upload_template(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Download Excel template file for pregnant women bulk upload
    Uses names instead of IDs for better usability
    Features:
    - Sheet 1: Template with human-readable headers
    - Sheet 2: Instructions with database reference lists
    """
    from fastapi.responses import StreamingResponse
    import pandas as pd
    import io
    from models import District, Block, Ward, SubCentre
    
    # Get sample data from database for current user's context
    district = None
    blocks = []
    wards = []
    sub_centres = []
    
    # Fetch real data based on user role
    if current_user.role == "district":
        district = db.query(District).filter(District.id == current_user.district_id).first()
        blocks = db.query(Block).filter(Block.district_id == current_user.district_id).limit(2).all()
        if blocks:
            wards = db.query(Ward).filter(Ward.block_id == blocks[0].id).limit(2).all()
            sub_centres = db.query(SubCentre).filter(SubCentre.block_id == blocks[0].id).limit(2).all()
    
    elif current_user.role == "block":
        block = db.query(Block).filter(Block.id == current_user.block_id).first()
        if block:
            district = db.query(District).filter(District.id == block.district_id).first()
            blocks = [block]
            wards = db.query(Ward).filter(Ward.block_id == block.id).limit(2).all()
            sub_centres = db.query(SubCentre).filter(SubCentre.block_id == block.id).limit(2).all()
    
    # ==================== SHEET 1: TEMPLATE DATA ====================
    # Create sample data with DB field names (will be renamed to human-readable headers)
    template_data = {
        "full_name": ["Jane Doe", "Mary Smith"],
        "mobile_number": ["9876543210", "9876543211"],
        # "aadhaar_number": ["123456789012", ""],
        "abha_id": ["12-3456-7890-1234", "12-3456-7890-1235"],
        "rch_id": ["RCH123456", "RCH123457"],
        "husband_name": ["John Doe", "John Smith"],
        "age": [25, 28],
        "district_name": [district.name if district else "Nashik", district.name if district else "Nashik"],
        "block_name": [blocks[0].name if blocks else "North Block", blocks[1].name if len(blocks) > 1 else blocks[0].name if blocks else "South Block"],
        "ward_name": [wards[0].name if wards else "Village A", wards[1].name if len(wards) > 1 else wards[0].name if wards else "Village B"],
        "sub_centre_name": [sub_centres[0].name if sub_centres else "PHC - Primary Health Centre 1", sub_centres[1].name if len(sub_centres) > 1 else sub_centres[0].name if sub_centres else "PHC - Primary Health Centre 2"],
        "date_of_birth": ["1999-01-01", "1996-05-15"],
        "lmp_date": ["2024-01-01", "2024-01-15"],
        "edd_date": ["2024-10-08", "2024-10-22"],
        "gravida": [2, 1],
        "para": [1, 0],
        "blood_group": ["O+", "A+"],
        "address": ["123 Main Street, Village A", "456 Oak Avenue, Village B"],
        "hpr_id": ["HPR123456", "HPR123457"]
    }
    
    # Create DataFrame
    df = pd.DataFrame(template_data)
    
    # ==================== MAPPING: DB FIELD → HUMAN-READABLE HEADER ====================
    column_mapping = {
        "full_name": "Full Name",
        "mobile_number": "Mobile Number",
        # "aadhaar_number": "Aadhaar Number",
        "abha_id": "ABHA ID",
        "rch_id": "RCH ID",
        "husband_name": "Husband Name",
        "age": "Age",
        "district_name": "District Name",
        "block_name": "Block Name",
        "ward_name": "Ward Name",
        "sub_centre_name": "Sub Centre Name",
        "date_of_birth": "Date of Birth",
        "lmp_date": "LMP Date",
        "edd_date": "EDD Date",
        "gravida": "Gravida",
        "para": "Para",
        "blood_group": "Blood Group",
        "address": "Address",
        "hpr_id": "HPR ID"
    }
    
    # Rename columns to human-readable headers
    df = df.rename(columns=column_mapping)
    
    # ==================== SHEET 2: INSTRUCTIONS WITH FIELD DESCRIPTIONS + DB REFERENCE LISTS ====================
    
    # Fetch all reference data from database based on user role
    all_districts = []
    all_blocks = []
    all_wards = []
    all_sub_centres = []
    
    if current_user.role == "district":
        # District user sees their district and all its blocks/wards/sub-centres
        all_districts = db.query(District).filter(District.id == current_user.district_id).all()
        all_blocks = db.query(Block).filter(Block.district_id == current_user.district_id).order_by(Block.name).all()
        
        # Get all wards in those blocks
        block_ids = [b.id for b in all_blocks]
        if block_ids:
            all_wards = db.query(Ward).filter(Ward.block_id.in_(block_ids)).order_by(Ward.name).all()
            all_sub_centres = db.query(SubCentre).filter(SubCentre.block_id.in_(block_ids)).order_by(SubCentre.name).all()
    
    elif current_user.role == "block":
        # Block user sees their block's district and block/wards/sub-centres
        block = db.query(Block).filter(Block.id == current_user.block_id).first()
        if block:
            all_districts = [db.query(District).filter(District.id == block.district_id).first()]
            all_blocks = [block]
            all_wards = db.query(Ward).filter(Ward.block_id == block.id).order_by(Ward.name).all()
            all_sub_centres = db.query(SubCentre).filter(SubCentre.block_id == block.id).order_by(SubCentre.name).all()
    
    # ==================== CREATE EXCEL FILE ====================
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Sheet 1: Template with human-readable headers
        df.to_excel(writer, sheet_name='Pregnant Women Data', index=False)
        
        # ==================== SHEET 2: COMBINED INSTRUCTIONS ====================
        # Part 1: Field Instructions Table
        field_instructions = pd.DataFrame({
            "Field Name": [
                "Full Name", "Mobile Number", # "Aadhaar Number", 
                "ABHA ID", "RCH ID",
                "Husband Name", "Age", "District Name", "Block Name", "Ward Name",
                "Sub Centre Name", "Date of Birth", "LMP Date", "EDD Date",
                "Gravida", "Para", "Blood Group", "Address", "HPR ID"
            ],
            "Required": [
                "Yes", "Yes", # "No", 
                "No", "No",
                "No", "No", "Yes", "Yes", "Yes",
                "Yes", "No", "No", "No",
                "No", "No", "No", "No", "No"
            ],
            "Description": [
                "Full name of pregnant woman",
                "10-digit mobile number",
                # "12-digit Aadhaar number (optional)",
                "ABHA ID in format XX-XXXX-XXXX-XXXX (optional)",
                "RCH ID (optional)",
                "Husband's full name",
                "Age in years",
                "District name (must exist in system - see reference list below)",
                "Block name (must exist in system - see reference list below)",
                "Ward/Village name - MANDATORY (must exist in system - see reference list below)",
                "Sub-Centre name - MANDATORY (must exist in system - see reference list below)",
                "Date of birth (YYYY-MM-DD format)",
                "Last Menstrual Period date (YYYY-MM-DD)",
                "Expected Delivery Date (YYYY-MM-DD)",
                "Number of pregnancies including current",
                "Number of live births",
                "Blood group (A+, A-, B+, B-, AB+, AB-, O+, O-)",
                "Full residential address",
                "Health Provider Registry ID (optional)"
            ],
            "Example": [
                "Sunita Devi", "9876543210", "12-3456-7890-1234", "RCH123456",
                "Ram Kumar", "25", district.name if district else "Nashik",
                blocks[0].name if blocks else "North Block",
                wards[0].name if wards else "Village A",
                sub_centres[0].name if sub_centres else "PHC - Primary Health Centre 1",
                "1999-01-01", "2024-01-01", "2024-10-08",
                "2", "1", "O+", "House No 123, Village A", "HPR123456"
            ]
        })
        
        # Write field instructions to Instructions sheet
        field_instructions.to_excel(writer, sheet_name='Instructions', index=False, startrow=0)
        
        # Part 2: Database Reference Lists
        # Create reference lists
        district_names = [d.name for d in all_districts if d]
        block_names = [b.name for b in all_blocks if b]
        ward_names = [w.name for w in all_wards if w]
        sub_centre_names = [sc.name for sc in all_sub_centres if sc]
        
        # Pad lists to same length for DataFrame (use empty strings for shorter lists)
        max_length = max(len(district_names), len(block_names), len(ward_names), len(sub_centre_names), 1)
        
        district_names += [""] * (max_length - len(district_names))
        block_names += [""] * (max_length - len(block_names))
        ward_names += [""] * (max_length - len(ward_names))
        sub_centre_names += [""] * (max_length - len(sub_centre_names))
        
        # Create reference lists DataFrame
        reference_lists = pd.DataFrame({
            "District Names": district_names,
            "Block Names": block_names,
            "Ward Names": ward_names,
            "Sub Centre Names": sub_centre_names
        })
        
        # Get the worksheet to append reference lists
        worksheet = writer.sheets['Instructions']
        
        # Calculate starting row for reference lists (after field instructions + 2 blank rows)
        start_row = len(field_instructions) + 3
        
        # Add section header for reference lists
        worksheet.cell(row=start_row, column=1, value="DATABASE REFERENCE LISTS (Copy exact names from below):")
        
        # Write reference lists below field instructions
        for r_idx, row in enumerate(reference_lists.itertuples(index=False), start=start_row + 2):
            for c_idx, value in enumerate(row, start=1):
                worksheet.cell(row=r_idx, column=c_idx, value=value)
        
        # Write column headers for reference lists
        ref_headers = ["District Names", "Block Names", "Ward Names", "Sub Centre Names"]
        for c_idx, header in enumerate(ref_headers, start=1):
            cell = worksheet.cell(row=start_row + 1, column=c_idx, value=header)
            # Make headers bold
            from openpyxl.styles import Font
            cell.font = Font(bold=True)
        
        # ==================== END SHEET 2 COMBINED INSTRUCTIONS ====================
        
        # Auto-adjust column widths for all sheets
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
    
    output.seek(0)
    
    # Return as downloadable file
    return StreamingResponse(
        io.BytesIO(output.read()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=pregnant_women_bulk_upload_template.xlsx"}
    )

@router.post("/forgot-password")
async def forgot_password(
    request: ForgotPasswordRequest,
    http_request: Request,
    db: Session = Depends(get_db)
):
    """
    Request password reset OTP via email
    """
    from models import PasswordResetToken
    from email_utils import generate_otp, get_otp_expiry, send_otp_email
    from datetime import datetime, timedelta
    
    # Find user by username or email
    user = db.query(User).filter(
        (User.username == request.identifier) | 
        (User.email == request.identifier)
    ).first()
    
    # Always return success message for security (don't reveal if user exists)
    success_response = {
        "message": "If the account exists, an OTP has been sent to your registered email address",
        "expires_in_minutes": int(os.getenv("OTP_EXPIRY_MINUTES", "15"))
    }
    
    if not user or not user.email:
        return success_response
    
    # Check rate limiting (max 3 requests per hour)
    one_hour_ago = datetime.now() - timedelta(hours=1)
    recent_requests = db.query(PasswordResetToken).filter(
        PasswordResetToken.user_id == user.id,
        PasswordResetToken.created_at > one_hour_ago
    ).count()
    
    if recent_requests >= 3:
        return success_response  # Don't reveal rate limiting
    
    # Invalidate existing tokens
    db.query(PasswordResetToken).filter(
        PasswordResetToken.user_id == user.id,
        PasswordResetToken.is_used == False
    ).update({"is_used": True})
    
    # Generate OTP and create token
    otp = generate_otp()
    expires_at = get_otp_expiry()
    
    reset_token = PasswordResetToken(
        user_id=user.id,
        token=otp,
        expires_at=expires_at
    )
    
    db.add(reset_token)
    db.commit()
    
    # Send OTP email
    email_sent = await send_otp_email(user.email, otp, user.full_name)
    
    if not email_sent:
        # If email fails, invalidate the token
        reset_token.is_used = True
        db.commit()
    
    return success_response

@router.post("/reset-password")
async def reset_password(
    request: ResetPasswordRequest,
    http_request: Request,
    db: Session = Depends(get_db)
):
    """
    Reset password using OTP
    """
    from models import PasswordResetToken, LoginAttempt
    from datetime import datetime
    
    # Find user
    user = db.query(User).filter(
        (User.username == request.identifier) | 
        (User.email == request.identifier)
    ).first()
    
    if not user:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid credentials"
        )
    
    # Find valid tokenen
    reset_token = db.query(PasswordResetToken).filter(
        PasswordResetToken.user_id == user.id,
        PasswordResetToken.token == request.otp,
        PasswordResetToken.is_used == False,
        PasswordResetToken.expires_at > datetime.now()
    ).first()
    
    if not reset_token:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid or expired OTP"
        )
    
    # Update password
    user.password_hash = get_password_hash(request.new_password)
    user.updated_at = datetime.now()
    
    # Mark token as used
    reset_token.is_used = True
    
    # Invalidate all other tokens for this user
    db.query(PasswordResetToken).filter(
        PasswordResetToken.user_id == user.id,
        PasswordResetToken.is_used == False
    ).update({"is_used": True})
    
    # Clear failed login attempts (unlock account)
    db.query(LoginAttempt).filter(
        LoginAttempt.username == user.username,
        LoginAttempt.success == False
    ).delete()
    
    db.commit()
    
    # Audit log: password reset
    ip_address, user_agent = get_client_info(http_request)
    log_password_reset(db, user.id, ip_address, user_agent)
    
    return {
        "message": "Password reset successfully. You can now login with your new password"
    }

@router.post("/change-password")
async def change_password(
    old_password: str,
    new_password: str,
    request: Request,
    current_user: User = Depends(get_current_active_user),
    db: Session = Depends(get_db)
):
    """
    Change user password
    """
    from auth import verify_password
    
    if not verify_password(old_password, current_user.password_hash):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Incorrect old password"
        )
    
    # Validate new password strength
    if len(new_password) < 8:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must be at least 8 characters long"
        )
    if not any(c.isupper() for c in new_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must contain at least one uppercase letter"
        )
    if not any(c.islower() for c in new_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must contain at least one lowercase letter"
        )
    if not any(c.isdigit() for c in new_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must contain at least one number"
        )
    if not any(c in '!@#$%^&*()_+-=[]{}|;:,.<>?' for c in new_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Password must contain at least one special character"
        )
    
    current_user.password_hash = get_password_hash(new_password)
    db.commit()
    
    # Audit log: password change
    ip_address, user_agent = get_client_info(request)
    log_password_change(db, current_user.id, ip_address, user_agent)
    
    # Send notification
    NotificationService.create_notification(
        db=db,
        user_ids=[current_user.id],
        title="Password Changed",
        message="Your password has been changed successfully. If you did not make this change, please contact administrator immediately.",
        notification_type="password_changed",
        category="user_management",
        priority="high",
        reference_id=current_user.id,
        reference_type="user",
        action_url="/profile",
        metadata={"username": current_user.username}
    )
    
    return {"message": "Password changed successfully"}

@router.post("/logout")
async def logout(
    request: Request,
    current_user: User = Depends(get_current_active_user),
    db: Session = Depends(get_db)
):
    """
    Logout endpoint to record logout event
    """
    ip_address, user_agent = get_client_info(request)
    log_logout(db, current_user.id, ip_address, user_agent)
    return {"message": "Logged out successfully"}
