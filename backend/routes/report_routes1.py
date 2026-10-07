from fastapi import APIRouter, Depends, HTTPException, status, BackgroundTasks
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func, and_, or_, case, desc, distinct
from typing import List, Optional, Dict, Any
from datetime import datetime, date, timedelta
import pandas as pd
import io
import os
import hashlib
import json
from functools import lru_cache

from database import get_db
from models import (
    User, PregnantWoman, ANCVisit, USGAppointment, Grievance,
    District, Block, Ward, SubCentre, USGCentre, SMSLog,
    DeliveryReferral, DeliveryOutcome, DeliveryPoint, ECGReport
)
from auth import get_current_active_user

router = APIRouter(prefix="/reports", tags=["Reports"])

# Simple in-memory cache for reports
report_cache: Dict[str, Dict[str, Any]] = {}
CACHE_TTL = 300  # 5 minutes

def get_cache_key(endpoint: str, params: dict) -> str:
    """Generate cache key from endpoint and parameters"""
    param_str = json.dumps(params, sort_keys=True, default=str)
    return hashlib.md5(f"{endpoint}:{param_str}".encode()).hexdigest()

def get_cached_report(cache_key: str) -> Optional[Dict[str, Any]]:
    """Get cached report if valid"""
    if cache_key in report_cache:
        cached = report_cache[cache_key]
        if datetime.now().timestamp() - cached['timestamp'] < CACHE_TTL:
            return cached['data']
        else:
            del report_cache[cache_key]
    return None

def set_cached_report(cache_key: str, data: Dict[str, Any]):
    """Cache report data"""
    report_cache[cache_key] = {
        'data': data,
        'timestamp': datetime.now().timestamp()
    }

def generate_visualization_config(data: list, chart_type: str, x_field: str, y_field: str, title: str) -> dict:
    """Generate chart configuration for frontend"""
    return {
        "type": chart_type,
        "title": title,
        "x_axis": x_field,
        "y_axis": y_field,
        "data": data
    }

# Column name mappings: db field -> friendly display name
COLUMN_LABELS = {
    # District Summary
    "total_pregnant_women": "Total Pregnant Women",
    "high_risk_cases": "High Risk Cases",
    "high_risk_percentage": "High Risk (%)",
    "self_registered": "Self Registered",
    # Block Wise Data
    "block_name": "Block Name",
    "total_women": "Total Women",
    "high_risk": "High Risk",
    "anc_visits": "ANC Visits",
    "usg_referrals": "USG Referrals",
    # USG Statistics
    "total_appointments": "Total Appointments",
    "completed_appointments": "Completed Appointments",
    "emergency_appointments": "Emergency Appointments",
    "completion_rate": "Completion Rate (%)",
    # Grievance Statistics
    "total_grievances": "Total Grievances",
    "resolved_grievances": "Resolved Grievances",
    "pending_grievances": "Pending Grievances",
    "escalated_grievances": "Escalated Grievances",
    "resolution_rate": "Resolution Rate (%)",
    "average_resolution_time_days": "Avg Resolution Time (Days)",
    # Ward Wise Data
    "ward_name": "Ward Name",
    "total_pregnant_women": "Total Pregnant Women",
    "high_risk_cases": "High Risk Cases",
    "approved_registrations": "Approved Registrations",
    "self_registrations": "Self Registrations",
    "total_anc_visits": "Total ANC Visits",
    "emergency_visits": "Emergency Visits",
    # Block Summary
    "total_wards": "Total Wards",
    "total_high_risk": "Total High Risk",
    # High Risk Tracking
    "id": "ID",
    "full_name": "Full Name",
    "mobile_number": "Mobile Number",
    "age": "Age",
    "risk_factors": "Risk Factors",
    "edd_date": "Expected Delivery Date",
    "last_anc_visit": "Last ANC Visit",
    "total_anc_visits": "Total ANC Visits",
    "usg_scans_completed": "USG Scans Completed",
    "block_name": "Block Name",
    "sub_centre_name": "Sub Centre Name",
    # ANC Compliance
    "expected_visits": "Expected Visits",
    "actual_visits": "Actual Visits",
    "compliance_status": "Compliance Status",
    "last_visit_date": "Last Visit Date",
    "days_since_last_visit": "Days Since Last Visit",
    "next_visit_due": "Next Visit Due",
    # USG Status
    "pregnant_woman_name": "Patient Name",
    "usg_centre_name": "USG Centre",
    "scheduled_date": "Scheduled Date",
    "status": "Status",
    "appointment_type": "Appointment Type",
    "days_pending": "Days Pending",
    "is_overdue": "Overdue",
    "completed_date": "Completed Date",
    # User Activity
    "user_name": "User Name",
    "role": "Role",
    "registrations": "Registrations",
    "total_actions": "Total Actions",
    # Comparative Analysis
    "name": "Name",
    "usg_completed": "USG Completed",
    "performance_score": "Performance Score",
    # Delivery Reporting
    "delivery_referrals": "Delivery Referrals",
    "completed_deliveries": "Completed Deliveries",
    "total_referrals": "Total Referrals",
    "pending_referrals": "Pending Referrals",
    "accepted_referrals": "Accepted Referrals",
    "safe_delivery": "Safe Delivery",
    "live_birth": "Live Birth",
    "still_birth": "Still Birth",
    "infant_death": "Infant Death",
    "maternal_death": "Maternal Death",
    "adverse_outcome_rate": "Adverse Outcome Rate (%)",
    "dp_name": "Delivery Point",
    "acceptance_rate": "Acceptance Rate (%)",
    "re_referral_rate": "Re-referral Rate (%)",
    "avg_acceptance_time_hours": "Avg Acceptance Time (Hours)",
    "total_delivery_referrals": "Total Delivery Referrals",
    "total_completed_deliveries": "Total Completed Deliveries",
}


def rename_columns(data: list) -> list:
    """Rename dict keys using COLUMN_LABELS mapping"""
    return [
        {COLUMN_LABELS.get(k, k.replace('_', ' ').title()): v for k, v in row.items()}
        for row in data
    ]


# Helper function to generate Excel file
async def generate_excel_report(data: dict, filename: str) -> StreamingResponse:
    """Generate Excel file from data dictionary with friendly column names"""
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for sheet_name, sheet_data in data.items():
            friendly_data = rename_columns(sheet_data) if sheet_data else sheet_data
            df = pd.DataFrame(friendly_data)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    output.seek(0)

    return StreamingResponse(
        io.BytesIO(output.read()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# Helper function to generate CSV file
async def generate_csv_report(data: list, filename: str) -> StreamingResponse:
    """Generate CSV file from data list with friendly column names"""
    df = pd.DataFrame(rename_columns(data) if data else data)
    output = io.StringIO()
    df.to_csv(output, index=False)
    output.seek(0)
    
    return StreamingResponse(
        io.StringIO(output.getvalue()),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

# ─── Delivery Helper Functions ───────────────────────────────────────────────

def _build_delivery_role_filter(query, current_user, referral_model=True):
    """
    Apply role-based filter to a DeliveryReferral or DeliveryOutcome query.
    Both models have pregnant_woman_id; geographic scope goes through PregnantWoman.
    For dp role, filter directly on dp_id — no PregnantWoman join needed.
    Returns the filtered query.
    """
    role = current_user.role
    if role == "district":
        query = query.join(PregnantWoman,
            (DeliveryReferral.pregnant_woman_id == PregnantWoman.id)
            if referral_model else
            (DeliveryOutcome.pregnant_woman_id == PregnantWoman.id)
        ).filter(PregnantWoman.district_id == current_user.district_id)
    elif role == "block":
        query = query.join(PregnantWoman,
            (DeliveryReferral.pregnant_woman_id == PregnantWoman.id)
            if referral_model else
            (DeliveryOutcome.pregnant_woman_id == PregnantWoman.id)
        ).filter(PregnantWoman.block_id == current_user.block_id)
    elif role == "sub_centre":
        query = query.join(PregnantWoman,
            (DeliveryReferral.pregnant_woman_id == PregnantWoman.id)
            if referral_model else
            (DeliveryOutcome.pregnant_woman_id == PregnantWoman.id)
        ).filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
    elif role == "dp":
        # Direct filter — no join required
        if referral_model:
            query = query.filter(DeliveryReferral.dp_id == current_user.dp_id)
        else:
            query = query.filter(DeliveryOutcome.dp_id == current_user.dp_id)
    return query


def get_delivery_summary_data(
    db: Session,
    current_user,
    start_date: date,
    end_date: date,
    block_id: Optional[int] = None
) -> dict:
    """
    Aggregated delivery summary for a given user scope and date range.
    Returns a plain dict — safe to call from any endpoint.
    All counts default to 0 on empty data.
    """
    try:
        # ── Referral counts by status ──────────────────────────────────────
        ref_q = db.query(
            DeliveryReferral.status,
            func.count(DeliveryReferral.id).label("cnt")
        )
        ref_q = _build_delivery_role_filter(ref_q, current_user, referral_model=True)
        ref_q = ref_q.filter(DeliveryReferral.created_at.between(start_date, end_date))
        if block_id and current_user.role == "district":
            # PregnantWoman already joined by helper
            ref_q = ref_q.filter(PregnantWoman.block_id == block_id)
        status_rows = ref_q.group_by(DeliveryReferral.status).all()
        # Normalize: MySQL may return enum name (PENDING) or value (pending) — always use lowercase value
        status_dict = {}
        for s, c in status_rows:
            key = s.value.lower() if hasattr(s, 'value') else str(s).split('.')[-1].lower()
            status_dict[key] = c

        # ── Outcome counts by delivery_type ────────────────────────────────
        out_q = db.query(
            DeliveryOutcome.delivery_type,
            func.count(DeliveryOutcome.id).label("cnt")
        )
        out_q = _build_delivery_role_filter(out_q, current_user, referral_model=False)
        out_q = out_q.filter(
            DeliveryOutcome.delivery_date >= start_date,
            DeliveryOutcome.delivery_date < end_date + timedelta(days=1)
        )
        if block_id and current_user.role == "district":
            out_q = out_q.filter(PregnantWoman.block_id == block_id)
        outcome_rows = out_q.group_by(DeliveryOutcome.delivery_type).all()
        # Normalize: MySQL may return enum name (SAFE_DELIVERY) or value (safe_delivery)
        outcome_dict = {}
        for dt, c in outcome_rows:
            key = dt.value.lower() if hasattr(dt, 'value') else str(dt).split('.')[-1].lower()
            outcome_dict[key] = c

        total_outcomes = sum(outcome_dict.values())
        adverse = (
            outcome_dict.get("still_birth", 0)
            + outcome_dict.get("infant_death", 0)
            + outcome_dict.get("maternal_death", 0)
        )
        adverse_rate = round((adverse / total_outcomes * 100) if total_outcomes > 0 else 0, 2)

        return {
            "total_referrals": sum(status_dict.values()),
            "pending": status_dict.get("pending", 0),
            "accepted": status_dict.get("accepted", 0),
            "re_referred": status_dict.get("re_referred", 0),
            "completed": status_dict.get("completed", 0),
            "outcomes": outcome_dict,
            "total_outcomes": total_outcomes,
            "adverse_rate": adverse_rate,
        }
    except Exception:
        return {
            "total_referrals": 0, "pending": 0, "accepted": 0,
            "re_referred": 0, "completed": 0,
            "outcomes": {}, "total_outcomes": 0, "adverse_rate": 0.0,
        }


def get_delivery_counts_by_ward(
    db: Session,
    ward_id: int,
    start_date: date,
    end_date: date
) -> tuple:
    """
    Returns (referral_count, completed_delivery_count) for a single ward.
    Uses aggregated SQL — no Python loops.
    Safe: returns (0, 0) on any error.
    """
    try:
        referral_count = (
            db.query(func.count(DeliveryReferral.id))
            .join(PregnantWoman, DeliveryReferral.pregnant_woman_id == PregnantWoman.id)
            .filter(
                PregnantWoman.ward_id == ward_id,
                DeliveryReferral.created_at.between(start_date, end_date)
            )
            .scalar() or 0
        )
        outcome_count = (
            db.query(func.count(DeliveryOutcome.id))
            .join(PregnantWoman, DeliveryOutcome.pregnant_woman_id == PregnantWoman.id)
            .filter(
                PregnantWoman.ward_id == ward_id,
                DeliveryOutcome.delivery_date >= start_date,
                DeliveryOutcome.delivery_date < end_date + timedelta(days=1)
            )
            .scalar() or 0
        )
        return (referral_count, outcome_count)
    except Exception:
        return (0, 0)


# ─── End Delivery Helper Functions ───────────────────────────────────────────


# DISTRICT USER REPORTS
@router.get("/district/performance")
async def get_district_performance_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    block_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Get district performance report (District users only)"""
    if current_user.role != "district":
        raise HTTPException(status_code=403, detail="District user access required")
    
    # Set default date range if not provided
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Base query for district data with date filter
    query = db.query(PregnantWoman).filter(
        PregnantWoman.district_id == current_user.district_id,
        PregnantWoman.created_at.between(start_date, end_date)
    )
    
    if block_id:
        query = query.filter(PregnantWoman.block_id == block_id)
    
    # Get district statistics
    total_pregnant_women = query.count()
    high_risk_cases = query.filter(PregnantWoman.is_high_risk == True).count()
    self_registered = query.filter(PregnantWoman.is_self_registered == True).count()
    
    # Get block-wise breakdown with date filter
    block_stats_query = db.query(
        Block.id.label('block_id'),
        Block.name.label('block_name'),
        func.count(PregnantWoman.id).label('total_women'),
        func.sum(case((PregnantWoman.is_high_risk == True, 1), else_=0)).label('high_risk'),
        func.sum(case((PregnantWoman.is_self_registered == True, 1), else_=0)).label('self_registered')
    ).join(PregnantWoman, Block.id == PregnantWoman.block_id)\
     .filter(
         Block.district_id == current_user.district_id,
         PregnantWoman.created_at.between(start_date, end_date)
     )
    
    if block_id:
        block_stats_query = block_stats_query.filter(Block.id == block_id)
    
    block_stats = block_stats_query.group_by(Block.id, Block.name).all()
    
    # Get ANC visits and USG referrals for each block
    block_wise_data = []
    for stat in block_stats:
        anc_visits = db.query(func.count(ANCVisit.id)).join(PregnantWoman)\
            .filter(
                PregnantWoman.block_id == stat.block_id,
                ANCVisit.visit_date.between(start_date, end_date)
            ).scalar() or 0
        
        usg_referrals = db.query(func.count(ANCVisit.id)).join(PregnantWoman)\
            .filter(
                PregnantWoman.block_id == stat.block_id,
                ANCVisit.referred_for_usg == True,
                ANCVisit.visit_date.between(start_date, end_date)
            ).scalar() or 0
        
        block_wise_data.append({
            "block_name": stat.block_name,
            "total_women": stat.total_women,
            "high_risk": stat.high_risk,
            "self_registered": stat.self_registered,
            "anc_visits": anc_visits,
            "usg_referrals": usg_referrals
        })
    
    # Get USG appointment statistics (date-filtered)
    usg_stats_query = db.query(
        func.count(USGAppointment.id).label('total_appointments'),
        func.sum(case((USGAppointment.status == 'completed', 1), else_=0)).label('completed'),
        func.sum(case((USGAppointment.appointment_type == 'emergency', 1), else_=0)).label('emergency')
    ).join(PregnantWoman, USGAppointment.pregnant_woman_id == PregnantWoman.id)\
     .filter(PregnantWoman.district_id == current_user.district_id)\
     .filter(USGAppointment.scheduled_date.between(start_date, end_date))
    
    if block_id:
        usg_stats_query = usg_stats_query.filter(PregnantWoman.block_id == block_id)
    
    usg_stats = usg_stats_query.first()
    
    # Get grievance statistics (date-filtered)
    grievance_stats_query = db.query(
        func.count(Grievance.id).label('total_grievances'),
        func.sum(case((Grievance.status == 'resolved', 1), else_=0)).label('resolved'),
        func.sum(case((Grievance.escalated_to_district == True, 1), else_=0)).label('escalated')
    ).filter(Grievance.district_id == current_user.district_id)\
     .filter(Grievance.created_at.between(start_date, end_date))
    
    if block_id:
        grievance_stats_query = grievance_stats_query.filter(Grievance.assigned_to_block == block_id)
    
    grievance_stats = grievance_stats_query.first()
    
    response = {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "district_summary": {
            "total_pregnant_women": total_pregnant_women,
            "high_risk_cases": high_risk_cases,
            "self_registered": self_registered,
            "high_risk_percentage": round((high_risk_cases / total_pregnant_women * 100) if total_pregnant_women > 0 else 0, 2)
        },
        "block_wise_data": block_wise_data,
        "usg_statistics": {
            "total_appointments": usg_stats.total_appointments or 0,
            "completed_appointments": usg_stats.completed or 0,
            "emergency_appointments": usg_stats.emergency or 0,
            "completion_rate": round((usg_stats.completed / usg_stats.total_appointments * 100) if usg_stats.total_appointments > 0 else 0, 2)
        },
        "grievance_statistics": {
            "total_grievances": grievance_stats.total_grievances or 0,
            "resolved_grievances": grievance_stats.resolved or 0,
            "escalated_grievances": grievance_stats.escalated or 0,
            "resolution_rate": round((grievance_stats.resolved / grievance_stats.total_grievances * 100) if grievance_stats.total_grievances > 0 else 0, 2)
        }
    }

    # --- delivery_statistics: NEW additive key (Step 5A) ---
    try:
        d = get_delivery_summary_data(db, current_user, start_date, end_date, block_id)
        response["delivery_statistics"] = {
            "total_referrals": d["total_referrals"],
            "completed_deliveries": d["completed"],
            "pending_referrals": d["pending"],
            "outcome_breakdown": {
                "safe_delivery":   d["outcomes"].get("safe_delivery", 0),
                "live_birth":      d["outcomes"].get("live_birth", 0),
                "still_birth":     d["outcomes"].get("still_birth", 0),
                "infant_death":    d["outcomes"].get("infant_death", 0),
                "maternal_death":  d["outcomes"].get("maternal_death", 0),
            },
            "adverse_outcome_rate": d["adverse_rate"],
        }
    except Exception:
        pass  # fail silently — never break existing response

    return response


def _get_block_detail_data(db: Session, block_id: int, start_date: date, end_date: date) -> dict:
    """Fetch ward summary and pregnant women list for a single block."""
    # --- Ward summary: all wards in block, with counts for women registered in date range ---
    wards = db.query(Ward).filter(Ward.block_id == block_id).all()

    ward_rows = []
    for ward in wards:
        total_women = db.query(func.count(PregnantWoman.id)).filter(
            PregnantWoman.ward_id == ward.id,
            PregnantWoman.created_at.between(start_date, end_date)
        ).scalar() or 0

        high_risk = db.query(func.count(PregnantWoman.id)).filter(
            PregnantWoman.ward_id == ward.id,
            PregnantWoman.is_high_risk == True,
            PregnantWoman.created_at.between(start_date, end_date)
        ).scalar() or 0

        anc_visits = db.query(func.count(ANCVisit.id)).join(PregnantWoman).filter(
            PregnantWoman.ward_id == ward.id,
            ANCVisit.visit_date.between(start_date, end_date)
        ).scalar() or 0

        usg_referrals = db.query(func.count(ANCVisit.id)).join(PregnantWoman).filter(
            PregnantWoman.ward_id == ward.id,
            ANCVisit.referred_for_usg == True,
            ANCVisit.visit_date.between(start_date, end_date)
        ).scalar() or 0

        emergency_visits = db.query(func.count(ANCVisit.id)).join(PregnantWoman).filter(
            PregnantWoman.ward_id == ward.id,
            ANCVisit.is_emergency == True,
            ANCVisit.visit_date.between(start_date, end_date)
        ).scalar() or 0

        ward_rows.append({
            "ward_name": ward.name,
            "total_women": total_women,
            "high_risk": high_risk,
            "anc_visits": anc_visits,
            "usg_referrals": usg_referrals,
            "emergency_visits": emergency_visits
        })

    # --- Pregnant women detail: ALL active women in block (not date-filtered) ---
    women = db.query(PregnantWoman).filter(
        PregnantWoman.block_id == block_id,
        PregnantWoman.is_active == True
    ).order_by(PregnantWoman.full_name).all()

    women_rows = []
    for pw in women:
        anc_count = db.query(func.count(ANCVisit.id)).filter(
            ANCVisit.pregnant_woman_id == pw.id
        ).scalar() or 0

        usg_count = db.query(func.count(USGAppointment.id)).filter(
            USGAppointment.pregnant_woman_id == pw.id,
            USGAppointment.status == "completed"
        ).scalar() or 0

        ward_name = None
        if pw.ward_id:
            ward = db.query(Ward).filter(Ward.id == pw.ward_id).first()
            ward_name = ward.name if ward else None

        sub_centre_name = None
        if pw.sub_centre_id:
            sc = db.query(SubCentre).filter(SubCentre.id == pw.sub_centre_id).first()
            sub_centre_name = sc.name if sc else None

        women_rows.append({
            "full_name": pw.full_name,
            "mobile_number": pw.mobile_number,
            "age": pw.age,
            "high_risk": "Yes" if pw.is_high_risk else "No",
            "risk_factors": pw.risk_factors or "",
            "lmp_date": pw.lmp_date,
            "edd_date": pw.edd_date,
            "anc_visits": anc_count,
            "usg_completed": usg_count,
            "registration_status": "Approved" if pw.registration_approved else "Pending",
            "ward_name": ward_name or "",
            "sub_centre_name": sub_centre_name or ""
        })

    return {"ward_rows": ward_rows, "women_rows": women_rows}


@router.get("/district/export/{format}")
async def export_district_report(
    format: str,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    block_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Export district report in specified format"""
    if current_user.role != "district":
        raise HTTPException(status_code=403, detail="District user access required")

    if format not in ["excel", "csv"]:
        raise HTTPException(status_code=400, detail="Format must be 'excel' or 'csv'")

    # Set default date range
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)

    # Get report data
    report_data = await get_district_performance_report(start_date, end_date, block_id, db, current_user)

    if format == "excel":
        output = io.BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # --- Sheet 1: District Summary ---
            pd.DataFrame(rename_columns([report_data["district_summary"]])).to_excel(
                writer, sheet_name="District Summary", index=False
            )

            # --- Sheet 2: Block Wise Data ---
            pd.DataFrame(rename_columns(report_data["block_wise_data"])).to_excel(
                writer, sheet_name="Block Wise Data", index=False
            )

            # --- Sheet 3: USG Statistics ---
            pd.DataFrame(rename_columns([report_data["usg_statistics"]])).to_excel(
                writer, sheet_name="USG Statistics", index=False
            )

            # --- Sheet 4: Grievance Statistics ---
            pd.DataFrame(rename_columns([report_data["grievance_statistics"]])).to_excel(
                writer, sheet_name="Grievance Statistics", index=False
            )

            # --- Sheet 5: Delivery Statistics (NEW — Step 5C) ---
            try:
                d = get_delivery_summary_data(db, current_user, start_date, end_date, block_id)

                # Row 1: district-level summary
                summary_row = [{
                    "total_referrals":      d["total_referrals"],
                    "pending_referrals":    d["pending"],
                    "accepted_referrals":   d["accepted"],
                    "completed_deliveries": d["completed"],
                    "safe_delivery":        d["outcomes"].get("safe_delivery", 0),
                    "live_birth":           d["outcomes"].get("live_birth", 0),
                    "still_birth":          d["outcomes"].get("still_birth", 0),
                    "infant_death":         d["outcomes"].get("infant_death", 0),
                    "maternal_death":       d["outcomes"].get("maternal_death", 0),
                    "adverse_outcome_rate": d["adverse_rate"],
                }]
                summary_df = pd.DataFrame(rename_columns(summary_row))
                summary_df.to_excel(writer, sheet_name="Delivery Statistics", index=False, startrow=0)

                # Rows below: one row per Delivery Point
                dp_q = db.query(DeliveryPoint).filter(
                    DeliveryPoint.is_active == True,
                    DeliveryPoint.district_id == current_user.district_id
                ).all()

                # Aggregated referral counts per DP
                ref_agg = db.query(
                    DeliveryReferral.dp_id,
                    DeliveryReferral.status,
                    func.count(DeliveryReferral.id).label("cnt")
                ).filter(
                    DeliveryReferral.created_at.between(start_date, end_date)
                ).group_by(DeliveryReferral.dp_id, DeliveryReferral.status).all()

                # Aggregated outcome counts per DP
                out_agg = db.query(
                    DeliveryOutcome.dp_id,
                    DeliveryOutcome.delivery_type,
                    func.count(DeliveryOutcome.id).label("cnt")
                ).filter(
                    DeliveryOutcome.delivery_date >= start_date,
                    DeliveryOutcome.delivery_date < end_date + timedelta(days=1)
                ).group_by(DeliveryOutcome.dp_id, DeliveryOutcome.delivery_type).all()

                ref_by_dp = {}
                for row in ref_agg:
                    if row.dp_id not in ref_by_dp:
                        ref_by_dp[row.dp_id] = {}
                    key = row.status.value.lower() if hasattr(row.status, 'value') else str(row.status).split('.')[-1].lower()
                    ref_by_dp[row.dp_id][key] = row.cnt

                out_by_dp = {}
                for row in out_agg:
                    if row.dp_id not in out_by_dp:
                        out_by_dp[row.dp_id] = {}
                    key = row.delivery_type.value.lower() if hasattr(row.delivery_type, 'value') else str(row.delivery_type).split('.')[-1].lower()
                    out_by_dp[row.dp_id][key] = row.cnt

                dp_rows = []
                for dp in dp_q:
                    ref = ref_by_dp.get(dp.id, {})
                    out = out_by_dp.get(dp.id, {})
                    total_ref = sum(ref.values())
                    dp_rows.append({
                        "Delivery Point": dp.name,
                        "Total Referrals": total_ref,
                        "Pending Referrals": ref.get("pending", 0),
                        "Accepted Referrals": ref.get("accepted", 0) + ref.get("completed", 0),
                        "Completed Deliveries": ref.get("completed", 0),
                        "Safe Delivery": out.get("safe_delivery", 0),
                        "Live Birth": out.get("live_birth", 0),
                        "Still Birth": out.get("still_birth", 0),
                        "Infant Death": out.get("infant_death", 0),
                        "Maternal Death": out.get("maternal_death", 0),
                    })

                if dp_rows:
                    from openpyxl.styles import Font as OFont
                    ws = writer.sheets["Delivery Statistics"]
                    # Gap row then DP header label
                    gap_row = len(summary_df) + 3  # +1 header +1 data +1 blank
                    ws.cell(row=gap_row, column=1, value="Delivery Point Wise Breakdown").font = OFont(bold=True)
                    pd.DataFrame(dp_rows).to_excel(
                        writer, sheet_name="Delivery Statistics", index=False, startrow=gap_row
                    )

            except Exception:
                pass  # fail silently — never break existing sheets

            # --- Per-block sheets: Ward Summary + Women Detail ---
            blocks_to_export = db.query(Block).filter(
                Block.district_id == current_user.district_id,
                Block.is_active == True
            )
            if block_id:
                blocks_to_export = blocks_to_export.filter(Block.id == block_id)
            blocks_to_export = blocks_to_export.all()

            from openpyxl.styles import Font

            for block in blocks_to_export:
                detail = _get_block_detail_data(db, block.id, start_date, end_date)

                # Excel sheet name max 31 chars
                sheet_name = block.name[:31]

                ward_rows = detail["ward_rows"]
                women_rows = detail["women_rows"]

                # Always write ward summary (even if empty, use placeholder)
                ward_df = pd.DataFrame(rename_columns(ward_rows)) if ward_rows else pd.DataFrame([{"Info": "No wards found for this block"}])
                women_df = pd.DataFrame(rename_columns(women_rows)) if women_rows else pd.DataFrame([{"Info": "No active pregnant women found for this block"}])

                # Section 1: Ward Summary header label at row 0, table at row 1
                ward_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)

                worksheet = writer.sheets[sheet_name]

                # Bold section label for ward summary
                ward_label = worksheet.cell(row=1, column=1, value="Ward Summary")
                ward_label.font = Font(bold=True)

                # Section 2: Women detail — starts after ward table + 2 gap rows
                women_start_row = 1 + len(ward_df) + 3

                # Bold section label for women detail
                women_label = worksheet.cell(row=women_start_row, column=1, value="Pregnant Women Detail")
                women_label.font = Font(bold=True)

                women_df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=women_start_row + 1)

        output.seek(0)
        filename = f"district_report_{start_date}_{end_date}.xlsx"
        return StreamingResponse(
            io.BytesIO(output.read()),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    elif format == "csv":
        filename = f"district_block_data_{start_date}_{end_date}.csv"
        return await generate_csv_report(report_data["block_wise_data"], filename)

# BLOCK USER REPORTS
@router.get("/block/ward-wise")
async def get_ward_wise_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    ward_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Get ward-wise report for block users"""
    if current_user.role != "block":
        raise HTTPException(status_code=403, detail="Block user access required")
    
    # Set default date range
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Base ward_stats query with date filter and ward_id filter
    ward_stats_query = db.query(
        Ward.id.label('ward_id'),
        func.coalesce(Ward.name, 'Rural Area (No Ward)').label('ward_name'),
        func.count(PregnantWoman.id).label('total_women'),
        func.sum(case((PregnantWoman.is_high_risk == True, 1), else_=0)).label('high_risk'),
        func.sum(case((PregnantWoman.registration_approved == True, 1), else_=0)).label('approved'),
        func.sum(case((PregnantWoman.is_self_registered == True, 1), else_=0)).label('self_registered')
    ).outerjoin(Ward, Ward.id == PregnantWoman.ward_id)\
     .filter(
         PregnantWoman.block_id == current_user.block_id,
         PregnantWoman.created_at.between(start_date, end_date)
     )

    if ward_id:
        ward_stats_query = ward_stats_query.filter(PregnantWoman.ward_id == ward_id)

    ward_stats = ward_stats_query.group_by(Ward.id, Ward.name).all()

    # Get ANC visit statistics by ward with date filter and ward_id filter
    anc_stats_query = db.query(
        Ward.id.label('ward_id'),
        func.coalesce(Ward.name, 'Rural Area (No Ward)').label('ward_name'),
        func.count(ANCVisit.id).label('total_visits'),
        func.sum(case((ANCVisit.referred_for_usg == True, 1), else_=0)).label('usg_referrals'),
        func.sum(case((ANCVisit.is_emergency == True, 1), else_=0)).label('emergency_visits')
    ).join(PregnantWoman, ANCVisit.pregnant_woman_id == PregnantWoman.id)\
     .outerjoin(Ward, PregnantWoman.ward_id == Ward.id)\
     .filter(
         PregnantWoman.block_id == current_user.block_id,
         ANCVisit.visit_date.between(start_date, end_date)
     )

    if ward_id:
        anc_stats_query = anc_stats_query.filter(PregnantWoman.ward_id == ward_id)

    anc_stats = anc_stats_query.group_by(Ward.id, Ward.name).all()
    
    # Combine ward statistics — match by ward_id (not name, to avoid duplicates)
    ward_data = []
    for ward_stat in ward_stats:
        anc_data = next((anc for anc in anc_stats if anc.ward_id == ward_stat.ward_id), None)

        # Step 5B: delivery counts per ward (safe, fail silently)
        try:
            dr_count, cd_count = get_delivery_counts_by_ward(
                db, ward_stat.ward_id, start_date, end_date
            )
        except Exception:
            dr_count, cd_count = 0, 0

        ward_data.append({
            "ward_name": ward_stat.ward_name,
            "total_pregnant_women": ward_stat.total_women,
            "high_risk_cases": ward_stat.high_risk,
            "approved_registrations": ward_stat.approved,
            "self_registrations": ward_stat.self_registered,
            "total_anc_visits": anc_data.total_visits if anc_data else 0,
            "usg_referrals": anc_data.usg_referrals if anc_data else 0,
            "emergency_visits": anc_data.emergency_visits if anc_data else 0,
            "delivery_referrals": dr_count,
            "completed_deliveries": cd_count,
        })

    return {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "block_summary": {
            "total_wards": len(ward_data),
            "total_pregnant_women": sum(w["total_pregnant_women"] for w in ward_data),
            "total_high_risk": sum(w["high_risk_cases"] for w in ward_data),
            "total_anc_visits": sum(w["total_anc_visits"] for w in ward_data),
            "total_delivery_referrals": sum(w["delivery_referrals"] for w in ward_data),
            "total_completed_deliveries": sum(w["completed_deliveries"] for w in ward_data),
        },
        "ward_wise_data": ward_data
    }

@router.get("/block/export/{format}")
async def export_block_report(
    format: str,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    ward_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Export block ward-wise report"""
    if current_user.role != "block":
        raise HTTPException(status_code=403, detail="Block user access required")
    
    if format not in ["excel", "csv"]:
        raise HTTPException(status_code=400, detail="Format must be 'excel' or 'csv'")
    
    # Get report data
    report_data = await get_ward_wise_report(start_date, end_date, ward_id, db, current_user)
    
    if format == "excel":
        # Step 5D: block_summary already has total_delivery_referrals and
        # total_completed_deliveries from get_ward_wise_report.
        # ward_wise_data already has delivery_referrals and completed_deliveries.
        # rename_columns inside generate_excel_report handles friendly labels.
        excel_data = {
            "Block Summary": [report_data["block_summary"]],
            "Ward Wise Data": report_data["ward_wise_data"]
        }
        filename = f"block_ward_report_{start_date or 'all'}_{end_date or 'all'}.xlsx"
        return await generate_excel_report(excel_data, filename)

    elif format == "csv":
        filename = f"ward_wise_data_{start_date or 'all'}_{end_date or 'all'}.csv"
        return await generate_csv_report(report_data["ward_wise_data"], filename)

# ============================================
# PHASE 1: ESSENTIAL REPORTS
# ============================================

# 1. HIGH-RISK PREGNANCY TRACKING REPORT
@router.get("/high-risk-tracking")
async def get_high_risk_tracking_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    page: int = 1,
    page_size: int = 50,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Track all high-risk pregnancies with detailed information"""
    # Check cache
    cache_key = get_cache_key("high-risk-tracking", {
        "user_id": current_user.id,
        "start_date": start_date,
        "end_date": end_date,
        "page": page
    })
    cached = get_cached_report(cache_key)
    if cached:
        return cached
    
    # Set default date range
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=90)
    
    # Base query with role-based filtering and eager loading
    query = db.query(PregnantWoman).options(
        joinedload(PregnantWoman.sub_centre)
    ).filter(
        PregnantWoman.is_high_risk == True,
        PregnantWoman.is_active == True
    )
    
    if current_user.role == "district":
        query = query.filter(PregnantWoman.district_id == current_user.district_id)
    elif current_user.role == "block":
        query = query.filter(PregnantWoman.block_id == current_user.block_id)
    elif current_user.role == "sub_centre":
        query = query.filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
    
    # Apply date filter
    query = query.filter(PregnantWoman.created_at.between(start_date, end_date))
    
    # Get total count
    total_count = query.count()
    
    # Pagination
    offset = (page - 1) * page_size
    high_risk_cases = query.order_by(desc(PregnantWoman.created_at)).offset(offset).limit(page_size).all()
    
    # Get risk factor breakdown
    risk_factors_count = {}
    for case in query.all():
        if case.risk_factors:
            factors = case.risk_factors.split(',')
            for factor in factors:
                factor = factor.strip()
                risk_factors_count[factor] = risk_factors_count.get(factor, 0) + 1
    
    # Get follow-up status (ANC visits)
    cases_with_followup = []
    for case in high_risk_cases:
        last_visit = db.query(ANCVisit).filter(
            ANCVisit.pregnant_woman_id == case.id
        ).order_by(desc(ANCVisit.visit_date)).first()
        
        usg_count = db.query(USGAppointment).filter(
            USGAppointment.pregnant_woman_id == case.id,
            USGAppointment.status == "completed"
        ).count()
        
        # Get block name if block_id exists
        block_name = None
        if case.block_id:
            block = db.query(Block).filter(Block.id == case.block_id).first()
            block_name = block.name if block else None
        
        cases_with_followup.append({
            "id": case.id,
            "full_name": case.full_name,
            "mobile_number": case.mobile_number,
            "age": case.age,
            "risk_factors": case.risk_factors,
            "edd_date": case.edd_date,
            "last_anc_visit": last_visit.visit_date if last_visit else None,
            "total_anc_visits": db.query(ANCVisit).filter(ANCVisit.pregnant_woman_id == case.id).count(),
            "usg_scans_completed": usg_count,
            "block_name": block_name,
            "sub_centre_name": case.sub_centre.name if case.sub_centre else None
        })
    
    result = {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "summary": {
            "total_high_risk_cases": total_count,
            "cases_with_recent_followup": sum(1 for c in cases_with_followup if c['last_anc_visit'] and (date.today() - c['last_anc_visit']).days <= 30),
            "cases_needing_attention": sum(1 for c in cases_with_followup if not c['last_anc_visit'] or (date.today() - c['last_anc_visit']).days > 30)
        },
        "risk_factors_breakdown": risk_factors_count,
        "cases": cases_with_followup,
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_pages": (total_count + page_size - 1) // page_size,
            "total_count": total_count
        },
        "visualization": generate_visualization_config(
            [{"factor": k, "count": v} for k, v in risk_factors_count.items()],
            "bar",
            "factor",
            "count",
            "Risk Factors Distribution"
        )
    }
    
    set_cached_report(cache_key, result)
    return result

# 2. ANC VISIT COMPLIANCE REPORT
@router.get("/anc-compliance")
async def get_anc_compliance_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    page: int = 1,
    page_size: int = 50,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Track ANC visit compliance and identify missed visits"""
    cache_key = get_cache_key("anc-compliance", {
        "user_id": current_user.id,
        "start_date": start_date,
        "end_date": end_date,
        "page": page
    })
    cached = get_cached_report(cache_key)
    if cached:
        return cached
    
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=90)

    # Base query with eager loading
    query = db.query(PregnantWoman).options(
        joinedload(PregnantWoman.sub_centre)
    ).filter(
        PregnantWoman.is_active == True,
        PregnantWoman.edd_date >= date.today(),        # only active pregnancies
        PregnantWoman.created_at >= start_date,        # registered within date range
        PregnantWoman.created_at <= end_date
    )
    
    if current_user.role == "district":
        query = query.filter(PregnantWoman.district_id == current_user.district_id)
    elif current_user.role == "block":
        query = query.filter(PregnantWoman.block_id == current_user.block_id)
    elif current_user.role == "sub_centre":
        query = query.filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
    
    total_count = query.count()
    offset = (page - 1) * page_size
    pregnant_women = query.offset(offset).limit(page_size).all()
    
    compliance_data = []
    compliant_count = 0
    non_compliant_count = 0
    
    for pw in pregnant_women:
        # Calculate expected visits based on pregnancy stage
        if pw.lmp_date:
            weeks_pregnant = (date.today() - pw.lmp_date).days // 7
            if weeks_pregnant < 12:
                expected_visits = 1
            elif weeks_pregnant < 28:
                expected_visits = 2
            elif weeks_pregnant < 36:
                expected_visits = 3
            else:
                expected_visits = 4
        else:
            expected_visits = 1
        
        actual_visits = db.query(ANCVisit).filter(
            ANCVisit.pregnant_woman_id == pw.id,
            ANCVisit.visit_date >= start_date,
            ANCVisit.visit_date <= end_date
        ).count()
        
        last_visit = db.query(ANCVisit).filter(
            ANCVisit.pregnant_woman_id == pw.id,
            ANCVisit.visit_date >= start_date,
            ANCVisit.visit_date <= end_date
        ).order_by(desc(ANCVisit.visit_date)).first()
        
        is_compliant = actual_visits >= expected_visits
        if is_compliant:
            compliant_count += 1
        else:
            non_compliant_count += 1
        
        # Get block name if block_id exists
        block_name = None
        if pw.block_id:
            block = db.query(Block).filter(Block.id == pw.block_id).first()
            block_name = block.name if block else None
        
        compliance_data.append({
            "id": pw.id,
            "full_name": pw.full_name,
            "mobile_number": pw.mobile_number,
            "expected_visits": expected_visits,
            "actual_visits": actual_visits,
            "compliance_status": "Compliant" if is_compliant else "Non-Compliant",
            "last_visit_date": last_visit.visit_date if last_visit else None,
            "days_since_last_visit": (date.today() - last_visit.visit_date).days if last_visit else None,
            "next_visit_due": last_visit.next_visit_date if last_visit else None,
            "block_name": block_name,
            "sub_centre_name": pw.sub_centre.name if pw.sub_centre else None
        })
    
    result = {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "summary": {
            "total_active_pregnancies": total_count,
            "compliant": compliant_count,
            "non_compliant": non_compliant_count,
            "compliance_rate": round((compliant_count / total_count * 100) if total_count > 0 else 0, 2)
        },
        "compliance_data": compliance_data,
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_pages": (total_count + page_size - 1) // page_size,
            "total_count": total_count
        },
        "visualization": generate_visualization_config(
            [{"status": "Compliant", "count": compliant_count}, {"status": "Non-Compliant", "count": non_compliant_count}],
            "pie",
            "status",
            "count",
            "ANC Compliance Status"
        )
    }
    
    set_cached_report(cache_key, result)
    return result

# 3. USG APPOINTMENT STATUS REPORT
@router.get("/usg-status")
async def get_usg_status_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    status_filter: Optional[str] = None,
    page: int = 1,
    page_size: int = 50,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Track USG appointment status including pending and overdue"""
    cache_key = get_cache_key("usg-status", {
        "user_id": current_user.id,
        "start_date": start_date,
        "end_date": end_date,
        "status": status_filter,
        "page": page
    })
    cached = get_cached_report(cache_key)
    if cached:
        return cached
    
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Base query
    query = db.query(USGAppointment).join(PregnantWoman)
    
    if current_user.role == "district":
        query = query.filter(PregnantWoman.district_id == current_user.district_id)
    elif current_user.role == "block":
        query = query.filter(PregnantWoman.block_id == current_user.block_id)
    elif current_user.role == "sub_centre":
        query = query.filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
    elif current_user.role == "usg_centre":
        query = query.filter(USGAppointment.usg_centre_id == current_user.usg_centre_id)
    
    query = query.filter(USGAppointment.scheduled_date.between(start_date, end_date))
    
    if status_filter:
        query = query.filter(USGAppointment.status == status_filter)
    
    # Get status breakdown
    status_counts = db.query(
        USGAppointment.status,
        func.count(USGAppointment.id).label('count')
    ).join(PregnantWoman)
    
    if current_user.role == "district":
        status_counts = status_counts.filter(PregnantWoman.district_id == current_user.district_id)
    elif current_user.role == "block":
        status_counts = status_counts.filter(PregnantWoman.block_id == current_user.block_id)
    elif current_user.role == "sub_centre":
        status_counts = status_counts.filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
    
    status_counts = status_counts.filter(
        USGAppointment.scheduled_date.between(start_date, end_date)
    ).group_by(USGAppointment.status).all()
    
    # Get overdue appointments
    overdue_query = db.query(USGAppointment).join(PregnantWoman).filter(
        USGAppointment.scheduled_date < date.today(),
        USGAppointment.scheduled_date >= start_date,
        USGAppointment.status.in_(["scheduled", "accepted"])
    )
    
    if current_user.role == "district":
        overdue_query = overdue_query.filter(PregnantWoman.district_id == current_user.district_id)
    elif current_user.role == "block":
        overdue_query = overdue_query.filter(PregnantWoman.block_id == current_user.block_id)
    
    overdue_count = overdue_query.count()
    
    # Pagination
    total_count = query.count()
    offset = (page - 1) * page_size
    appointments = query.order_by(desc(USGAppointment.scheduled_date)).offset(offset).limit(page_size).all()
    
    appointments_data = []
    for apt in appointments:
        pw = db.query(PregnantWoman).filter(PregnantWoman.id == apt.pregnant_woman_id).first()
        usg_centre = db.query(USGCentre).filter(USGCentre.id == apt.usg_centre_id).first()
        
        days_pending = (date.today() - apt.scheduled_date.date()).days if apt.scheduled_date else 0
        is_overdue = apt.scheduled_date and apt.scheduled_date.date() < date.today() and apt.status in ["scheduled", "accepted"]
        
        appointments_data.append({
            "id": apt.id,
            "pregnant_woman_name": pw.full_name if pw else None,
            "mobile_number": pw.mobile_number if pw else None,
            "usg_centre_name": usg_centre.name if usg_centre else None,
            "scheduled_date": apt.scheduled_date,
            "status": apt.status,
            "appointment_type": apt.appointment_type,
            "days_pending": days_pending if apt.status in ["scheduled", "accepted"] else 0,
            "is_overdue": is_overdue,
            "completed_date": apt.completed_date
        })
    
    result = {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "summary": {
            "total_appointments": total_count,
            "status_breakdown": {status: count for status, count in status_counts},
            "overdue_appointments": overdue_count,
            "completion_rate": round(
                (dict(status_counts).get("completed", 0) / total_count * 100) if total_count > 0 else 0, 2
            )
        },
        "appointments": appointments_data,
        "pagination": {
            "page": page,
            "page_size": page_size,
            "total_pages": (total_count + page_size - 1) // page_size,
            "total_count": total_count
        },
        "visualization": generate_visualization_config(
            [{"status": status, "count": count} for status, count in status_counts],
            "bar",
            "status",
            "count",
            "USG Appointment Status Distribution"
        )
    }
    
    set_cached_report(cache_key, result)
    return result

# 4. GRIEVANCE RESOLUTION PERFORMANCE REPORT
@router.get("/grievance-performance")
async def get_grievance_performance_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Analyze grievance resolution performance"""
    cache_key = get_cache_key("grievance-performance", {
        "user_id": current_user.id,
        "start_date": start_date,
        "end_date": end_date
    })
    cached = get_cached_report(cache_key)
    if cached:
        return cached
    
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Base query
    query = db.query(Grievance).filter(
        Grievance.created_at.between(start_date, end_date)
    )
    
    if current_user.role == "district":
        query = query.filter(Grievance.district_id == current_user.district_id)
    elif current_user.role == "block":
        query = query.filter(Grievance.assigned_to_block == current_user.block_id)
    
    total_grievances = query.count()
    resolved_grievances = query.filter(Grievance.status == "resolved").all()
    escalated_grievances = query.filter(Grievance.escalated_to_district == True).count()
    
    # Calculate average resolution time
    resolution_times = []
    for grievance in resolved_grievances:
        if grievance.resolved_at:
            resolution_time = (grievance.resolved_at - grievance.created_at).days
            resolution_times.append(resolution_time)
    
    avg_resolution_time = sum(resolution_times) / len(resolution_times) if resolution_times else 0
    
    # Status breakdown
    status_breakdown = db.query(
        Grievance.status,
        func.count(Grievance.id).label('count')
    ).filter(Grievance.created_at.between(start_date, end_date))
    
    if current_user.role == "district":
        status_breakdown = status_breakdown.filter(Grievance.district_id == current_user.district_id)
    elif current_user.role == "block":
        status_breakdown = status_breakdown.filter(Grievance.assigned_to_block == current_user.block_id)
    
    status_breakdown = status_breakdown.group_by(Grievance.status).all()
    
    result = {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "summary": {
            "total_grievances": total_grievances,
            "resolved_grievances": len(resolved_grievances),
            "pending_grievances": total_grievances - len(resolved_grievances),
            "escalated_grievances": escalated_grievances,
            "resolution_rate": round((len(resolved_grievances) / total_grievances * 100) if total_grievances > 0 else 0, 2),
            "average_resolution_time_days": round(avg_resolution_time, 1)
        },
        "status_breakdown": {status: count for status, count in status_breakdown},
        "resolution_time_distribution": {
            "0-3_days": sum(1 for t in resolution_times if t <= 3),
            "4-7_days": sum(1 for t in resolution_times if 3 < t <= 7),
            "8-14_days": sum(1 for t in resolution_times if 7 < t <= 14),
            "15+_days": sum(1 for t in resolution_times if t > 14)
        },
        "visualization": generate_visualization_config(
            [{"status": status, "count": count} for status, count in status_breakdown],
            "pie",
            "status",
            "count",
            "Grievance Status Distribution"
        )
    }
    
    set_cached_report(cache_key, result)
    return result

# SUB-CENTRE USER REPORTS
@router.get("/sub-centre/activity")
async def get_sub_centre_activity_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Get sub-centre activity report"""
    if current_user.role != "sub_centre":
        raise HTTPException(status_code=403, detail="Sub-centre user access required")
    
    # Set default date range
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Get pregnant women under care
    pregnant_women = db.query(PregnantWoman).filter(
        PregnantWoman.sub_centre_id == current_user.sub_centre_id,
        PregnantWoman.is_active == True
    ).count()
    
    # Get ANC visits conducted
    anc_visits = db.query(ANCVisit).join(PregnantWoman)\
        .filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)\
        .filter(ANCVisit.visit_date.between(start_date, end_date)).count()
    
    # Get USG referrals made
    usg_referrals = db.query(ANCVisit).join(PregnantWoman)\
        .filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)\
        .filter(ANCVisit.referred_for_usg == True)\
        .filter(ANCVisit.visit_date.between(start_date, end_date)).count()
    
    return {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "activity_summary": {
            "pregnant_women_under_care": pregnant_women,
            "anc_visits_conducted": anc_visits,
            "usg_referrals_made": usg_referrals,
            "average_visits_per_day": round(anc_visits / ((end_date - start_date).days + 1), 2)
        }
    }

# USG CENTRE USER REPORTS
@router.get("/usg-centre/appointments")
async def get_usg_centre_appointment_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Get USG centre appointment report"""
    if current_user.role != "usg_centre":
        raise HTTPException(status_code=403, detail="USG centre user access required")
    
    # Set default date range
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Get appointment statistics
    total_appointments = db.query(USGAppointment).filter(
        USGAppointment.usg_centre_id == current_user.usg_centre_id,
        USGAppointment.scheduled_date.between(start_date, end_date)
    ).count()
    
    completed_appointments = db.query(USGAppointment).filter(
        USGAppointment.usg_centre_id == current_user.usg_centre_id,
        USGAppointment.status == "completed",
        USGAppointment.scheduled_date.between(start_date, end_date)
    ).count()
    
    emergency_appointments = db.query(USGAppointment).filter(
        USGAppointment.usg_centre_id == current_user.usg_centre_id,
        USGAppointment.appointment_type == "emergency",
        USGAppointment.scheduled_date.between(start_date, end_date)
    ).count()
    
    return {
        "report_period": {
            "start_date": start_date,
            "end_date": end_date
        },
        "appointment_summary": {
            "total_appointments": total_appointments,
            "completed_appointments": completed_appointments,
            "emergency_appointments": emergency_appointments,
            "completion_rate": round((completed_appointments / total_appointments * 100) if total_appointments > 0 else 0, 2),
            "average_appointments_per_day": round(total_appointments / ((end_date - start_date).days + 1), 2)
        }
    }


# ============================================
# PHASE 2: ADVANCED ANALYTICS
# ============================================

# 5. COMPARATIVE PERFORMANCE ANALYSIS
@router.get("/comparative-analysis")
async def get_comparative_analysis(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    comparison_type: str = "block",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Compare performance across blocks or sub-centres"""
    if current_user.role not in ["district", "block"]:
        raise HTTPException(status_code=403, detail="District or Block user access required")

    # Validate comparison_type vs role
    if comparison_type == "sub_centre" and current_user.role == "district":
        raise HTTPException(
            status_code=400,
            detail="District users can only use comparison_type=block. Use block user login for sub_centre comparison."
        )
    if comparison_type == "block" and current_user.role == "block":
        raise HTTPException(
            status_code=400,
            detail="Block users can only use comparison_type=sub_centre."
        )

    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)

    comparative_data = []

    # --- District user: compare blocks ---
    if comparison_type == "block" and current_user.role == "district":
        blocks = db.query(Block).filter(
            Block.district_id == current_user.district_id,
            Block.is_active == True
        ).all()

        for block in blocks:
            total_pw = db.query(PregnantWoman).filter(
                PregnantWoman.block_id == block.id,
                PregnantWoman.is_active == True
            ).count()

            high_risk = db.query(PregnantWoman).filter(
                PregnantWoman.block_id == block.id,
                PregnantWoman.is_high_risk == True
            ).count()

            anc_visits = db.query(ANCVisit).join(PregnantWoman).filter(
                PregnantWoman.block_id == block.id,
                ANCVisit.visit_date.between(start_date, end_date)
            ).count()

            usg_completed = db.query(USGAppointment).join(PregnantWoman).filter(
                PregnantWoman.block_id == block.id,
                USGAppointment.status == "completed",
                USGAppointment.scheduled_date.between(start_date, end_date)
            ).count()

            grievances_resolved = db.query(Grievance).filter(
                Grievance.assigned_to_block == block.id,
                Grievance.status == "resolved",
                Grievance.created_at.between(start_date, end_date)
            ).count()

            comparative_data.append({
                "id": block.id,
                "name": block.name,
                "total_pregnant_women": total_pw,
                "high_risk_cases": high_risk,
                "anc_visits": anc_visits,
                "usg_completed": usg_completed,
                "grievances_resolved": grievances_resolved,
                "performance_score": round(
                    (anc_visits / total_pw if total_pw > 0 else 0) * 50 +
                    (usg_completed / total_pw if total_pw > 0 else 0) * 50
                , 2)
            })

    # --- Block user: compare sub-centres ---
    elif comparison_type == "sub_centre" and current_user.role == "block":
        sub_centres = db.query(SubCentre).filter(
            SubCentre.block_id == current_user.block_id,
            SubCentre.is_active == True
        ).all()

        for sc in sub_centres:
            total_pw = db.query(PregnantWoman).filter(
                PregnantWoman.sub_centre_id == sc.id,
                PregnantWoman.is_active == True
            ).count()

            high_risk = db.query(PregnantWoman).filter(
                PregnantWoman.sub_centre_id == sc.id,
                PregnantWoman.is_high_risk == True
            ).count()

            anc_visits = db.query(ANCVisit).join(PregnantWoman).filter(
                PregnantWoman.sub_centre_id == sc.id,
                ANCVisit.visit_date.between(start_date, end_date)
            ).count()

            usg_completed = db.query(USGAppointment).join(PregnantWoman).filter(
                PregnantWoman.sub_centre_id == sc.id,
                USGAppointment.status == "completed",
                USGAppointment.scheduled_date.between(start_date, end_date)
            ).count()

            usg_referrals = db.query(ANCVisit).join(PregnantWoman).filter(
                PregnantWoman.sub_centre_id == sc.id,
                ANCVisit.referred_for_usg == True,
                ANCVisit.visit_date.between(start_date, end_date)
            ).count()

            comparative_data.append({
                "id": sc.id,
                "name": sc.name,
                "total_pregnant_women": total_pw,
                "high_risk_cases": high_risk,
                "anc_visits": anc_visits,
                "usg_referrals": usg_referrals,
                "usg_completed": usg_completed,
                "performance_score": round(
                    (anc_visits / total_pw if total_pw > 0 else 0) * 50 +
                    (usg_completed / total_pw if total_pw > 0 else 0) * 50
                , 2)
            })

    comparative_data.sort(key=lambda x: x.get('performance_score', 0), reverse=True)

    return {
        "report_period": {"start_date": start_date, "end_date": end_date},
        "comparison_type": comparison_type,
        "total_entities": len(comparative_data),
        "data": comparative_data,
        "best_performer": comparative_data[0] if comparative_data else None,
        "worst_performer": comparative_data[-1] if len(comparative_data) > 1 else None,
        "visualization": generate_visualization_config(
            comparative_data, "bar", "name", "performance_score", "Comparative Performance"
        )
    }

# 6. TREND ANALYSIS
@router.get("/trends")
async def get_trend_analysis(
    months: int = 6,
    metric: str = "registrations",
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Analyze trends over time"""
    end_date = date.today()
    start_date = end_date - timedelta(days=months * 30)
    
    trend_data = []
    current_date = start_date
    
    while current_date <= end_date:
        month_start = current_date.replace(day=1)
        month_end = (month_start.replace(month=month_start.month % 12 + 1, day=1) if month_start.month < 12 else month_start.replace(year=month_start.year + 1, month=1, day=1)) - timedelta(days=1)

        count = 0  # default for every iteration

        if metric == "registrations":
            query = db.query(func.count(PregnantWoman.id)).filter(PregnantWoman.created_at.between(month_start, month_end))
            if current_user.role == "district":
                query = query.filter(PregnantWoman.district_id == current_user.district_id)
            elif current_user.role == "block":
                query = query.filter(PregnantWoman.block_id == current_user.block_id)
            elif current_user.role == "sub_centre":
                query = query.filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
            count = query.scalar() or 0

        elif metric == "usg_appointments":
            query = db.query(func.count(USGAppointment.id)).filter(USGAppointment.scheduled_date.between(month_start, month_end))
            if current_user.role == "district":
                query = query.join(PregnantWoman).filter(PregnantWoman.district_id == current_user.district_id)
            elif current_user.role == "block":
                query = query.join(PregnantWoman).filter(PregnantWoman.block_id == current_user.block_id)
            elif current_user.role == "usg_centre":
                query = query.filter(USGAppointment.usg_centre_id == current_user.usg_centre_id)
            count = query.scalar() or 0

        elif metric == "anc_visits":
            query = db.query(func.count(ANCVisit.id)).filter(ANCVisit.visit_date.between(month_start, month_end))
            if current_user.role == "district":
                query = query.join(PregnantWoman).filter(PregnantWoman.district_id == current_user.district_id)
            elif current_user.role == "block":
                query = query.join(PregnantWoman).filter(PregnantWoman.block_id == current_user.block_id)
            elif current_user.role == "sub_centre":
                query = query.join(PregnantWoman).filter(PregnantWoman.sub_centre_id == current_user.sub_centre_id)
            count = query.scalar() or 0

        elif metric == "grievances":
            query = db.query(func.count(Grievance.id)).filter(Grievance.created_at.between(month_start, month_end))
            if current_user.role == "district":
                query = query.filter(Grievance.district_id == current_user.district_id)
            elif current_user.role == "block":
                query = query.filter(Grievance.assigned_to_block == current_user.block_id)
            count = query.scalar() or 0

        trend_data.append({"month": month_start.strftime("%Y-%m"), "count": count})
        current_date = month_end + timedelta(days=1)
    
    return {
        "metric": metric,
        "trend_data": trend_data,
        "summary": {"total": sum(d['count'] for d in trend_data), "average_per_month": round(sum(d['count'] for d in trend_data) / len(trend_data), 2) if trend_data else 0},
        "visualization": generate_visualization_config(trend_data, "line", "month", "count", f"{metric.title()} Trend")
    }

# ============================================
# PHASE 3: OPERATIONAL REPORTS
# ============================================

# 8. USER ACTIVITY REPORT
@router.get("/user-activity")
async def get_user_activity_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Track user activity and productivity"""
    if current_user.role not in ["district", "block"]:
        raise HTTPException(status_code=403, detail="District or Block user access required")
    
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Get users in jurisdiction
    query = db.query(User).filter(User.is_active == True)
    if current_user.role == "district":
        query = query.filter(User.district_id == current_user.district_id)
    elif current_user.role == "block":
        query = query.filter(User.block_id == current_user.block_id)
    
    users = query.all()
    user_activity = []
    
    for user in users:
        # Count registrations
        registrations = db.query(PregnantWoman).filter(PregnantWoman.registered_by == user.id, PregnantWoman.created_at.between(start_date, end_date)).count()
        # Count ANC visits
        anc_visits = db.query(ANCVisit).filter(ANCVisit.attended_by == user.id, ANCVisit.visit_date.between(start_date, end_date)).count()
        
        user_activity.append({
            "user_name": user.full_name,
            "role": user.role,
            "registrations": registrations,
            "anc_visits": anc_visits,
            "total_actions": registrations + anc_visits
        })
    
    user_activity.sort(key=lambda x: x['total_actions'], reverse=True)
    
    return {
        "report_period": {"start_date": start_date, "end_date": end_date},
        "user_activity": user_activity,
        "summary": {"total_users": len(users), "active_users": sum(1 for u in user_activity if u['total_actions'] > 0)},
        "visualization": generate_visualization_config(user_activity[:10], "bar", "user_name", "total_actions", "Top 10 Active Users")
    }

# 9. SMS DELIVERY REPORT
@router.get("/sms-delivery")
async def get_sms_delivery_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Track SMS delivery performance"""
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    
    # Get SMS logs
    sms_logs = db.query(SMSLog).filter(SMSLog.sent_at.between(start_date, end_date)).all()
    
    total_sms = len(sms_logs)
    delivered = sum(1 for log in sms_logs if log.delivery_status == "delivered")
    failed = sum(1 for log in sms_logs if log.delivery_status == "failed")
    
    # Message type breakdown
    type_breakdown = {}
    for log in sms_logs:
        type_breakdown[log.message_type] = type_breakdown.get(log.message_type, 0) + 1
    
    return {
        "report_period": {"start_date": start_date, "end_date": end_date},
        "summary": {
            "total_sms": total_sms,
            "delivered": delivered,
            "failed": failed,
            "delivery_rate": round((delivered / total_sms * 100) if total_sms > 0 else 0, 2)
        },
        "type_breakdown": type_breakdown,
        "visualization": generate_visualization_config([{"type": k, "count": v} for k, v in type_breakdown.items()], "pie", "type", "count", "SMS by Type")
    }

# ============================================
# DELIVERY REPORTS (Step 4)
# ============================================

@router.get("/delivery/summary")
async def get_delivery_summary_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Delivery referral and outcome summary for current user scope"""
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)

    d = get_delivery_summary_data(db, current_user, start_date, end_date)

    return {
        "report_period": {"start_date": start_date, "end_date": end_date},
        "referral_summary": {
            "total_referrals": d["total_referrals"],
            "pending": d["pending"],
            "accepted": d["accepted"],
            "re_referred": d["re_referred"],
            "completed": d["completed"],
        },
        "outcome_summary": {
            "total_outcomes": d["total_outcomes"],
            "by_type": d["outcomes"],
            "adverse_outcome_rate": d["adverse_rate"],
        },
        "visualization": generate_visualization_config(
            [{"type": k, "count": v} for k, v in d["outcomes"].items()],
            "donut", "type", "count", "Delivery Outcomes"
        ),
    }


@router.get("/delivery/outcome-breakdown")
async def get_delivery_outcome_breakdown(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Delivery outcome breakdown with percentages and adverse summary"""
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)

    try:
        out_q = db.query(
            DeliveryOutcome.delivery_type,
            func.count(DeliveryOutcome.id).label("cnt")
        )
        out_q = _build_delivery_role_filter(out_q, current_user, referral_model=False)
        out_q = out_q.filter(
            DeliveryOutcome.delivery_date >= start_date,
            DeliveryOutcome.delivery_date < end_date + timedelta(days=1)
        )
        rows = out_q.group_by(DeliveryOutcome.delivery_type).all()
        outcome_dict = {}
        for dt, c in rows:
            key = dt.value.lower() if hasattr(dt, 'value') else str(dt).split('.')[-1].lower()
            outcome_dict[key] = c
        total = sum(outcome_dict.values())

        breakdown = [
            {
                "delivery_type": dt,
                "count": cnt,
                "percentage": round((cnt / total * 100) if total > 0 else 0, 2),
            }
            for dt, cnt in outcome_dict.items()
        ]

        adverse_types = {"still_birth", "infant_death", "maternal_death"}
        adverse_total = sum(outcome_dict.get(t, 0) for t in adverse_types)

        return {
            "report_period": {"start_date": start_date, "end_date": end_date},
            "total_outcomes": total,
            "outcome_breakdown": breakdown,
            "adverse_summary": {
                "still_birth":    outcome_dict.get("still_birth", 0),
                "infant_death":   outcome_dict.get("infant_death", 0),
                "maternal_death": outcome_dict.get("maternal_death", 0),
                "total_adverse":  adverse_total,
                "adverse_rate":   round((adverse_total / total * 100) if total > 0 else 0, 2),
            },
            "visualization": generate_visualization_config(
                breakdown, "pie", "delivery_type", "count", "Outcome Breakdown"
            ),
        }
    except Exception:
        return {
            "report_period": {"start_date": start_date, "end_date": end_date},
            "total_outcomes": 0,
            "outcome_breakdown": [],
            "adverse_summary": {
                "still_birth": 0, "infant_death": 0,
                "maternal_death": 0, "total_adverse": 0, "adverse_rate": 0.0,
            },
            "visualization": generate_visualization_config([], "pie", "delivery_type", "count", "Outcome Breakdown"),
        }


@router.get("/delivery/point-performance")
async def get_delivery_point_performance(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Per-DP performance metrics. District and block roles only."""
    if current_user.role not in ["district", "block"]:
        raise HTTPException(status_code=403, detail="District or Block user access required")

    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)

    try:
        # Scope DPs to user jurisdiction
        dp_q = db.query(DeliveryPoint).filter(DeliveryPoint.is_active == True)
        if current_user.role == "district":
            dp_q = dp_q.filter(DeliveryPoint.district_id == current_user.district_id)
        elif current_user.role == "block":
            dp_q = dp_q.filter(DeliveryPoint.block_id == current_user.block_id)
        dps = dp_q.all()

        from sqlalchemy import text as sa_text

        # Single aggregated query: referral counts by dp_id and status
        # avg_accept_hours: use TIMESTAMPDIFF via raw text expression for MariaDB compatibility
        ref_agg = db.query(
            DeliveryReferral.dp_id,
            DeliveryReferral.status,
            func.count(DeliveryReferral.id).label("cnt"),
            func.avg(
                sa_text("TIMESTAMPDIFF(HOUR, delivery_referrals.created_at, delivery_referrals.accepted_at)")
            ).label("avg_accept_hours")
        ).filter(
            DeliveryReferral.created_at.between(start_date, end_date)
        ).group_by(DeliveryReferral.dp_id, DeliveryReferral.status).all()

        # Single aggregated query: outcome counts by dp_id and type
        # Fix: DateTime column needs >= start and < end+1 day (not .between)
        out_agg = db.query(
            DeliveryOutcome.dp_id,
            DeliveryOutcome.delivery_type,
            func.count(DeliveryOutcome.id).label("cnt")
        ).filter(
            DeliveryOutcome.delivery_date >= start_date,
            DeliveryOutcome.delivery_date < end_date + timedelta(days=1)
        ).group_by(DeliveryOutcome.dp_id, DeliveryOutcome.delivery_type).all()

        # Index by dp_id
        ref_by_dp: Dict[int, dict] = {}
        for row in ref_agg:
            dp_id = row.dp_id
            if dp_id not in ref_by_dp:
                ref_by_dp[dp_id] = {"pending": 0, "accepted": 0, "re_referred": 0, "completed": 0, "avg_accept_hours": 0.0}
            key = row.status.value.lower() if hasattr(row.status, 'value') else str(row.status).split('.')[-1].lower()
            ref_by_dp[dp_id][key] = row.cnt
            if key == "accepted" and row.avg_accept_hours:
                ref_by_dp[dp_id]["avg_accept_hours"] = round(float(row.avg_accept_hours), 1)

        out_by_dp: Dict[int, dict] = {}
        for row in out_agg:
            dp_id = row.dp_id
            if dp_id not in out_by_dp:
                out_by_dp[dp_id] = {}
            key = row.delivery_type.value.lower() if hasattr(row.delivery_type, 'value') else str(row.delivery_type).split('.')[-1].lower()
            out_by_dp[dp_id][key] = row.cnt

        dp_performance = []
        for dp in dps:
            ref = ref_by_dp.get(dp.id, {})
            out = out_by_dp.get(dp.id, {})
            total_received = ref.get("pending", 0) + ref.get("accepted", 0) + ref.get("re_referred", 0) + ref.get("completed", 0)
            accepted = ref.get("accepted", 0) + ref.get("completed", 0)
            re_referred = ref.get("re_referred", 0)

            dp_performance.append({
                "dp_id": dp.id,
                "dp_name": dp.name,
                "total_referrals_received": total_received,
                "accepted": accepted,
                "re_referred": re_referred,
                "completed": ref.get("completed", 0),
                "acceptance_rate": round((accepted / total_received * 100) if total_received > 0 else 0, 2),
                "re_referral_rate": round((re_referred / total_received * 100) if total_received > 0 else 0, 2),
                "avg_acceptance_time_hours": ref.get("avg_accept_hours", 0.0),
                "outcomes": {
                    "safe_delivery":  out.get("safe_delivery", 0),
                    "live_birth":     out.get("live_birth", 0),
                    "still_birth":    out.get("still_birth", 0),
                    "infant_death":   out.get("infant_death", 0),
                    "maternal_death": out.get("maternal_death", 0),
                },
            })

        dp_performance.sort(key=lambda x: x["completed"], reverse=True)

        return {
            "report_period": {"start_date": start_date, "end_date": end_date},
            "dp_performance": dp_performance,
            "visualization": generate_visualization_config(
                dp_performance, "bar", "dp_name", "completed", "Delivery Point Performance"
            ),
        }
    except Exception as e:
        import logging
        logging.getLogger(__name__).error(f"delivery/point-performance error: {e}", exc_info=True)
        return {
            "report_period": {"start_date": start_date, "end_date": end_date},
            "dp_performance": [],
            "visualization": generate_visualization_config([], "bar", "dp_name", "completed", "Delivery Point Performance"),
        }


# CLEAR CACHE ENDPOINT
@router.post("/clear-cache")
async def clear_report_cache(current_user: User = Depends(get_current_active_user)):
    """Clear report cache (admin only)"""
    if current_user.role != "district":
        raise HTTPException(status_code=403, detail="District user access required")
    
    report_cache.clear()
    return {"message": "Report cache cleared successfully"}


# ECG SUMMARY REPORT
@router.get("/ecg-summary")
async def get_ecg_summary_report(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """ECG report summary — accessible to district, block, sub_centre, dp roles."""
    if current_user.role not in ("district", "block", "sub_centre", "dp"):
        raise HTTPException(status_code=403, detail="Not authorized")

    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)

    query = db.query(ECGReport).filter(
        ECGReport.ecg_date >= start_date,
        ECGReport.ecg_date <= end_date
    )

    if current_user.role == "dp":
        query = query.filter(ECGReport.dp_id == current_user.dp_id)
    elif current_user.role == "sub_centre":
        pw_ids = db.query(PregnantWoman.id).filter(
            PregnantWoman.sub_centre_id == current_user.sub_centre_id
        ).subquery()
        query = query.filter(ECGReport.pregnant_woman_id.in_(pw_ids))
    elif current_user.role == "block":
        pw_ids = db.query(PregnantWoman.id).filter(
            PregnantWoman.block_id == current_user.block_id
        ).subquery()
        query = query.filter(ECGReport.pregnant_woman_id.in_(pw_ids))
    elif current_user.role == "district":
        # Match via pregnant woman's district_id OR via the DP's district_id
        pw_ids = db.query(PregnantWoman.id).filter(
            PregnantWoman.district_id == current_user.district_id
        ).subquery()
        dp_ids = db.query(DeliveryPoint.id).filter(
            DeliveryPoint.district_id == current_user.district_id
        ).subquery()
        query = query.filter(
            (ECGReport.pregnant_woman_id.in_(pw_ids)) |
            (ECGReport.dp_id.in_(dp_ids))
        )

    total = query.count()
    normal = query.filter(ECGReport.result == "normal").count()
    abnormal = query.filter(ECGReport.result == "abnormal").count()

    # Per-DP breakdown (useful for district/block)
    dp_breakdown = []
    if current_user.role in ("district", "block"):
        dp_rows = db.query(
            ECGReport.dp_id,
            func.count(ECGReport.id).label("total"),
            func.sum(case((ECGReport.result == "normal", 1), else_=0)).label("normal"),
            func.sum(case((ECGReport.result == "abnormal", 1), else_=0)).label("abnormal")
        ).filter(
            ECGReport.ecg_date >= start_date,
            ECGReport.ecg_date <= end_date
        )
        if current_user.role == "block":
            dp_rows = dp_rows.join(PregnantWoman, ECGReport.pregnant_woman_id == PregnantWoman.id).filter(
                PregnantWoman.block_id == current_user.block_id
            )
        elif current_user.role == "district":
            dp_ids_district = db.query(DeliveryPoint.id).filter(
                DeliveryPoint.district_id == current_user.district_id
            ).subquery()
            dp_rows = dp_rows.filter(ECGReport.dp_id.in_(dp_ids_district))
        dp_rows = dp_rows.group_by(ECGReport.dp_id).all()
        for row in dp_rows:
            dp = db.query(DeliveryPoint).filter(DeliveryPoint.id == row.dp_id).first()
            dp_breakdown.append({
                "dp_id": row.dp_id,
                "dp_name": dp.name if dp else None,
                "total": row.total,
                "normal": row.normal,
                "abnormal": row.abnormal,
            })

    return {
        "report_period": {"start_date": start_date, "end_date": end_date},
        "summary": {
            "total_ecg_reports": total,
            "normal": normal,
            "abnormal": abnormal,
            "abnormal_rate": round((abnormal / total * 100) if total > 0 else 0, 2),
        },
        "dp_breakdown": dp_breakdown,
    }
